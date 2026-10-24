"""Build the Second Look 2:30 pitch deck (docs/presentation/second-look-pitch.pptx).

Run: python build_deck.py <output.pptx>   (needs python-pptx)
Every number on a slide carries its source; speaker notes hold the full timed script.
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

INK = RGBColor(0x14, 0x1B, 0x2D)
MUTED = RGBColor(0x5B, 0x63, 0x75)
ACCENT = RGBColor(0x0F, 0x6E, 0x56)   # deep green
WARN = RGBColor(0xB4, 0x4D, 0x12)     # labels for fictional data
FONT = 'Arial'
DEMO_URL = 'https://x1cloud.tail6fbca0.ts.net/'

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, runs, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT):
    """runs: str or list of paragraphs; each paragraph is str or (str, size, color, bold)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    paras = runs if isinstance(runs, list) else [runs]
    for i, p in enumerate(paras):
        t, s, c, b = (p, size, color, bold) if isinstance(p, str) else p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.space_after = Pt(6)
        r = para.add_run()
        r.text = t
        r.font.name, r.font.size, r.font.color.rgb, r.font.bold = FONT, Pt(s), c, b
    return box


def slide(title, notes, kicker=None):
    s = prs.slides.add_slide(BLANK)
    if kicker:
        text(s, 0.8, 0.45, 11.7, 0.4, kicker, size=14, color=ACCENT, bold=True)
    if title:
        text(s, 0.8, 0.8, 11.7, 1.2, title, size=40, bold=True)
    s.notes_slide.notes_text_frame.text = notes
    return s


def source(s, line):
    text(s, 0.8, 6.85, 11.7, 0.4, line, size=11, color=MUTED)


# 1. Title ----------------------------------------------------------------
s = slide(None, '[0:00–0:05]\nHi, I\'m Ryan. This is Second Look.', kicker='DEAL FLOW · TRACK 4: FUNDS')
text(s, 0.8, 2.3, 11.7, 1.4, 'Second Look', size=66, bold=True)
text(s, 0.8, 3.7, 11.7, 1.0, 'Approve the deal you are actually getting.', size=30, color=MUTED)

# 2. Problem --------------------------------------------------------------
s = slide('LPs sell on an old report. The deal keeps moving.', (
    '[0:05–0:30]\n'
    'Investors selling a fund stake decide from a quarterly report that’s months old, while the companies move and the buyer revises terms. '
    'LPs sold about 125 billion dollars of fund stakes last year, and this year they’re averaging around 87 cents on the dollar, '
    'so one point of price is worth roughly 1.4 billion a year. And about 40 percent of these sellers are first-timers.\n\n'
    'TO VERIFY before presenting: $125B (Jefferies FY2025; Evercore says $120B), 87% of NAV (Jefferies H1 2026; bases may differ), '
    '~40% first-time sellers (Future Standard). The market doc says several of these came via blogs republishing the reports. '
    '$1.4B is our derivation: $125B / 0.87 ≈ $144B of NAV × 1%.'), kicker='THE PROBLEM')
stats = [('$125B', 'LP-led secondary volume, 2025'), ('87%', 'average LP pricing, % of NAV, H1 2026'),
         ('≈$1.4B', 'value of one point of price, per year'), ('~40%', 'of LP-led sellers are first-timers')]
for i, (big, small) in enumerate(stats):
    x = 0.8 + i * 3.0
    text(s, x, 2.6, 2.8, 1.0, big, size=48, color=ACCENT, bold=True)
    text(s, x, 3.7, 2.7, 1.2, small, size=17, color=INK)
source(s, 'Sources: Jefferies Global Secondary Market Review FY2025 and July 2026 (Evercore: $120B); Future Standard. '
          '$1.4B derived: $125B ÷ 0.87 ≈ $144B NAV × 1%.')

# 3. Existing solutions ---------------------------------------------------
s = slide('Today’s tools each cover part of it', (
    '[0:30–0:50]\n'
    'Marketplaces like Nasdaq Fund Secondaries and LP Desk are built to execute the trade. '
    'Monitoring tools like Chronograph, Cobalt and Landytech work from the same periodic manager reports. '
    'Advisors use email and spreadsheets, where answers scatter. '
    'And an AI chat reads documents but keeps no record of who approved which terms.\n\n'
    'If asked: the closest comparable is DiligenceVault, which offers versioned diligence Q&A and approval trails (per its own site). '
    'Our difference: the manager’s answer becomes an explicit term of this sale, and a changed term automatically retires the old approval.'),
    kicker='WHY EXISTING SOLUTIONS FALL SHORT')
rows = [('Secondary marketplaces', 'Nasdaq Fund Secondaries, LP Desk', 'Built to execute the trade, not to check the seller’s evidence'),
        ('LP monitoring', 'Chronograph, Cobalt, Landytech', 'Mainly built on periodic GP reports, not a specific sale’s terms'),
        ('Advisors, email, spreadsheets', '', 'Answers scatter; nothing flags a term change after approval'),
        ('General AI chat', '', 'Reads documents, but no shared record or versioned approvals')]
for i, (who, names, gap) in enumerate(rows):
    y = 2.2 + i * 1.1
    text(s, 0.8, y, 4.2, 1.0, [(who, 22, INK, True), (names, 14, MUTED, False)] if names else [(who, 22, INK, True)])
    text(s, 5.2, y + 0.05, 7.4, 1.0, gap, size=20, color=INK)
source(s, 'Closest comparable: DiligenceVault (versioned diligence Q&A, approval trails). Descriptions from vendors’ public positioning; not a hands-on evaluation.')

# 4. How it works (live demo) ---------------------------------------------
s = slide('One record, from evidence to the exact deal', (
    '[0:50–1:50] DEMO (about 60 seconds; click as you speak)\n'
    'The fund and deal are fictional; the game data is real. '
    'Someone offers me 700,000 dollars for my gaming-fund stake. '
    'I open Tiny Glade: real Steam reviews from before and after the last report, and Akta’s studio profile. '
        'Roundtable and Nicrron turn that evidence into a question for the manager; PRISM records the trace. '
    'I share it. As the manager: a 100,000 dollar distribution belongs to this stake. '
    'As the investor: 800,000 dollars gross. I approve version one. '
    'Now the buyer moves that distribution to themselves. Same 700,000 price, but I receive 700,000 gross, '
    'and my old approval no longer applies. I reload: the record is still here.\n\n'
    'Note: the hosted demo opens a saved review captured earlier today. Say “captured earlier today,” not “live.” '
    'If the live demo fails, use the next slide.'), kicker='OUR PRODUCT · DEMO')
steps = ['Real evidence\nSteam + Akta', 'Question to the GP\nRoundtable · Nicrron · PRISM', 'GP answers\n(fictional)',
         'LP approves\noffer v1', 'Buyer changes\na term (v2)', 'Old approval\nno longer applies']
for i, st in enumerate(steps):
    x = 0.8 + i * 2.05
    head, sub = st.split('\n')
    text(s, x, 2.6, 1.9, 0.5, f'{i + 1}', size=30, color=ACCENT, bold=True)
    text(s, x, 3.2, 1.9, 1.6, [(head, 18, INK, True), (sub, 14, MUTED, False)])
text(s, 0.8, 5.3, 11.7, 0.6, 'Every step is saved and survives a reload.', size=22, color=INK)
source(s, 'Real: Steam reviews and player snapshot, Akta company profiles. Fictional: the fund, its holdings, the offer and the GP’s answer. '
          'Steam data prompts questions; it is never a valuation.')

# 5. Same price, different deal (demo payoff / backup) --------------------
s = slide('Same price. Different deal.', (
    '[shown at about 1:40–1:50, or as backup if the live demo fails]\n'
    'Same 700,000 dollar headline, but moving one distribution changes what the seller receives by 100,000 dollars. '
    'The old approval stays on record, but it no longer applies to this version. '
    'These amounts are a fictional scenario, calculated by the server in whole cents, excluding tax, fees and timing.'),
    kicker='THE CONSEQUENCE')
cols = [('Offer v1', 'Seller keeps the $100,000 distribution', '$800,000', 'Approved'),
        ('Offer v2', 'Buyer takes the distribution', '$700,000', 'Old approval no longer applies')]
for i, (ver, term, amt, status) in enumerate(cols):
    x = 0.8 + i * 6.0
    text(s, x, 2.2, 5.6, 0.6, f'{ver} · $700,000 purchase price', size=20, color=MUTED, bold=True)
    text(s, x, 2.8, 5.6, 0.6, term, size=20)
    text(s, x, 3.5, 5.6, 1.2, amt, size=60, color=ACCENT if i == 0 else INK, bold=True)
    text(s, x, 4.8, 5.6, 0.6, status, size=20, color=WARN if i else ACCENT, bold=True)
text(s, 0.8, 5.7, 11.7, 0.6, 'Scenario gross receipts to the seller: −$100,000', size=22, bold=True)
source(s, 'FICTIONAL scenario: fund, offer and $100,000 distribution are illustrative. Server-calculated in integer cents; excludes tax, fees and timing.')

# 6. Market ---------------------------------------------------------------
s = slide('A small, high-value wedge in a large market', (
    '[1:50–2:20]\n'
    'Honestly sized: for the sale decision, our reachable market is an estimated 10 to 25 million dollars a year. '
    'But one point of price on a 75 million dollar stake is about 900,000 dollars, so a small fraction of a point pays for the tool, '
    'and we’ll test pricing above our 500-dollar-a-month hypothesis. '
    'Gaming funds are our beachhead because their holdings leave public evidence. '
    'The bigger opportunity is the LP’s ongoing evidence and approval record.\n\n'
    'Do NOT state billions as our market. $10–25M is a Second Look estimate built on assumptions (share of LP teams that see an offer, '
    'average deal size, price per offer). $75M is an illustrative stake (assumption): ≈$86M NAV at 87%, so one point ≈ $0.9M. '
    'No customer interviews yet: pricing is a hypothesis to test.'), kicker='IMPACT AND MARKET')
mk = [('$10–25M / yr', 'Reachable market for the sale decision', 'Second Look estimate'),
      ('≈$0.9M', 'One point of price on a $75M stake', 'Illustrative'),
      ('Gaming funds', 'First beachhead: holdings with public evidence', 'Not the TAM')]
for i, (big, small, tag) in enumerate(mk):
    x = 0.8 + i * 4.0
    text(s, x, 2.3, 3.8, 1.0, big, size=40, color=ACCENT, bold=True)
    text(s, x, 3.4, 3.6, 1.0, small, size=18)
    text(s, x, 4.3, 3.6, 0.5, tag.upper(), size=12, color=WARN, bold=True)
text(s, 0.8, 5.2, 11.7, 0.8, [('Next: the LP’s ongoing evidence and approval record across all fund interests.', 20, INK, True),
                              ('Pricing is a hypothesis to test: start at $500 per team per month; the value at stake suggests testing higher.', 16, MUTED, False)])
source(s, 'Second Look estimates (Oct 24, 2026) from stated assumptions; $75M stake is illustrative, ≈$86M NAV at Jefferies’ H1 2026 87%. $500/team/month is an untested hypothesis.')

# 7. Close ----------------------------------------------------------------
s = slide(None, (
    '[2:20–2:30]\n'
    'Investors approve fund sales on outdated information, and terms change under them. '
    'Second Look puts the new evidence to the manager and retires an approval the moment the terms change. '
    'We’re Mat, Ryan, Darren and Stratten. The demo link is on screen.\n\n'
    'Never show or say the access code; share it with judges privately.'), kicker='SECOND LOOK')
text(s, 0.8, 1.6, 11.7, 2.6, 'Investors approve fund sales on outdated information, and terms change under them. '
     'Second Look puts the new evidence to the manager and retires an approval the moment the terms change.', size=36, bold=True)
text(s, 0.8, 4.6, 11.7, 0.6, DEMO_URL, size=24, color=ACCENT, bold=True)
text(s, 0.8, 5.3, 11.7, 0.6, 'Mat Silverstein · Ryan · Darren · Stratten', size=20, color=MUTED)
source(s, 'Demo: fictional fund and deal; real Steam and Akta data. Team login required; ask us for access.')

prs.save(sys.argv[1] if len(sys.argv) > 1 else 'second-look-pitch.pptx')
